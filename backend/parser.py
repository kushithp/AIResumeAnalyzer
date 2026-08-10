import io
import re
from typing import Dict, List, Any
import pdfplumber
import docx
from utils import ALL_SKILLS, extract_contact_info


def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF file bytes using pdfplumber."""
    extracted_text = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                extracted_text.append(text)
    return "\n".join(extracted_text)


def parse_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX file bytes using python-docx."""
    doc = docx.Document(io.BytesIO(file_bytes))
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    full_text.append(cell.text.strip())
    return "\n".join(full_text)


def extract_skills_from_text(text: str) -> List[str]:
    """Scan resume text against standard skill database."""
    text_lower = text.lower()
    found_skills = set()
    
    for skill in ALL_SKILLS:
        # Match whole word boundaries
        pattern = r'\b' + re.escape(skill) + r'\b'
        if re.search(pattern, text_lower):
            found_skills.add(skill.title() if len(skill) > 3 else skill.upper())
            
    return sorted(list(found_skills))


def detect_sections(text: str) -> Dict[str, bool]:
    """Check for standard resume sections."""
    text_lower = text.lower()
    sections = {
        "Education": bool(re.search(r'\b(education|academic|university|degree|bachelor|master|gpa)\b', text_lower)),
        "Experience": bool(re.search(r'\b(experience|employment|work history|career|internship)\b', text_lower)),
        "Skills": bool(re.search(r'\b(skills|technical skills|technologies|proficiencies|competencies)\b', text_lower)),
        "Projects": bool(re.search(r'\b(projects|personal projects|portfolio|key projects)\b', text_lower)),
        "Certifications": bool(re.search(r'\b(certifications|certificates|licenses|courses)\b', text_lower))
    }
    return sections


def estimate_experience_level(text: str) -> str:
    """Estimate experience level based on year keywords and work patterns."""
    text_lower = text.lower()
    
    # Look for patterns like "5 years", "3+ years of experience"
    year_matches = re.findall(r'(\d+)\+?\s*years?\s*(of)?\s*(experience|work)?', text_lower)
    max_years = 0
    for match in year_matches:
        try:
            years = int(match[0])
            if years > max_years and years < 40:
                max_years = years
        except ValueError:
            pass

    if max_years >= 7 or "senior" in text_lower or "lead" in text_lower or "architect" in text_lower:
        return "Senior"
    elif max_years >= 3 or "mid-level" in text_lower or "software engineer ii" in text_lower:
        return "Mid-Level"
    elif max_years >= 1 or "junior" in text_lower or "associate" in text_lower:
        return "Junior"
    else:
        return "Fresher"


def parse_resume_file(filename: str, file_bytes: bytes) -> Dict[str, Any]:
    """Main parser interface for incoming resume files."""
    ext = filename.split('.')[-1].lower()
    
    if ext == 'pdf':
        text = parse_pdf(file_bytes)
    elif ext in ['docx', 'doc']:
        text = parse_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file format: .{ext}. Please upload a PDF or DOCX file.")
        
    if not text or len(text.strip()) < 20:
        raise ValueError("Could not extract readable text from the file. Please ensure it is not a scanned image PDF.")

    name, email, phone = extract_contact_info(text)
    skills = extract_skills_from_text(text)
    sections = detect_sections(text)
    exp_level = estimate_experience_level(text)

    return {
        "text": text,
        "candidate_name": name,
        "email": email,
        "phone": phone,
        "extracted_skills": skills,
        "sections": sections,
        "experience_level": exp_level
    }
